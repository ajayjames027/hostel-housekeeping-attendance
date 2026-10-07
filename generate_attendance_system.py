"""
Hostel Housekeeping Attendance & Remuneration Generator
Staff:
1. Ms. Arputhamani - MBA Hostel
2. Mr. R. Sakthivelan - Sports Hostel
3. Ms. Revathi - Toulouse Arena

Generates:
1. Excel Workbook (.xlsx) with formulas for Attendance Sign-sheet, Digital Log, and Remuneration/Payroll calculation.
2. Printable HTML Attendance Sign-in Register (Ready to print on A4 Landscape).
3. Printable HTML Remuneration Slip & Salary Register (Ready to print on A4 Portrait).
"""

import calendar
import datetime
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def generate_excel_system(month, year, hostel_title, staff_list, output_path):
    wb = openpyxl.Workbook()
    
    # Define styles
    font_title = Font(name="Calibri", size=16, bold=True, color="1F497D")
    font_subtitle = Font(name="Calibri", size=11, italic=True, color="595959")
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_bold = Font(name="Calibri", size=11, bold=True)
    font_regular = Font(name="Calibri", size=10)
    font_small = Font(name="Calibri", size=9)
    font_kpi = Font(name="Calibri", size=13, bold=True, color="1F497D")
    
    fill_navy = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    fill_steel = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
    fill_light_blue = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
    fill_weekend = PatternFill(start_color="F2DCDB", end_color="F2DCDB", fill_type="solid") # Soft red/peach for Sunday
    fill_summary = PatternFill(start_color="EBF1DE", end_color="EBF1DE", fill_type="solid") # Soft green
    
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    dark_border = Border(
        left=Side(style='thin', color='595959'),
        right=Side(style='thin', color='595959'),
        top=Side(style='thin', color='595959'),
        bottom=Side(style='thin', color='595959')
    )
    double_bottom_border = Border(
        top=Side(style='thin', color='595959'),
        bottom=Side(style='double', color='595959'),
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9')
    )

    num_days = calendar.monthrange(year, month)[1]
    month_name = calendar.month_name[month]

    # ==========================================
    # SHEET 1: Staff Master & Configuration
    # ==========================================
    ws_staff = wb.active
    ws_staff.title = "Staff_Master"
    ws_staff.views.sheetView[0].showGridLines = True

    ws_staff["A1"] = f"{hostel_title}"
    ws_staff["A1"].font = font_title
    ws_staff["A2"] = f"Housekeeping Staff Directory & Master Settings | {month_name} {year}"
    ws_staff["A2"].font = font_subtitle

    headers_staff = [
        "Staff ID", "Full Name", "Assigned Location / Hostel", "Designation",
        "Monthly Base Wage (₹)", "Daily Wage Rate (₹)", "Weekly Off", "Payment Mode", "Bank / UPI Details"
    ]
    
    for col_idx, h in enumerate(headers_staff, 1):
        cell = ws_staff.cell(row=4, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_navy
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    row_idx = 5
    for emp in staff_list:
        ws_staff.cell(row=row_idx, column=1, value=emp["id"]).alignment = Alignment(horizontal="center")
        ws_staff.cell(row=row_idx, column=2, value=emp["name"]).font = font_bold
        ws_staff.cell(row=row_idx, column=3, value=emp["location"]).font = font_bold
        ws_staff.cell(row=row_idx, column=4, value=emp["role"])
        ws_staff.cell(row=row_idx, column=5, value=emp["monthly_wage"]).number_format = "₹#,##0.00"
        # Daily wage formula = Monthly Wage / Total days in month
        ws_staff.cell(row=row_idx, column=6, value=f"=E{row_idx}/{num_days}").number_format = "₹#,##0.00"
        ws_staff.cell(row=row_idx, column=7, value=emp["weekly_off"]).alignment = Alignment(horizontal="center")
        ws_staff.cell(row=row_idx, column=8, value=emp["payment_mode"]).alignment = Alignment(horizontal="center")
        ws_staff.cell(row=row_idx, column=9, value=emp["account_details"])
        
        for c in range(1, 10):
            ws_staff.cell(row=row_idx, column=c).border = thin_border
        row_idx += 1

    # Add quick notes & instructions
    ws_staff.cell(row=row_idx + 2, column=1, value="System Instructions:").font = font_bold
    notes = [
        "1. Printable_Sign_Sheet: Print at the start of the month and display at respective hostel/arena housekeeping desks.",
        "2. Attendance_Log: Daily attendance (P = Present, HD = Half Day, WO = Weekly Off, PL = Paid Leave, A = Absent).",
        "3. Monthly_Remuneration: Calculates gross and net wages automatically based on payable days.",
        "4. Ready for wage disbursement with signature acknowledgement."
    ]
    for i, note in enumerate(notes, row_idx + 3):
        ws_staff.cell(row=i, column=1, value=note).font = font_regular

    # Auto-adjust column widths
    for col in ws_staff.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_staff.column_dimensions[col_letter].width = max(max_len + 4, 12)
    ws_staff.column_dimensions["A"].width = 12
    ws_staff.column_dimensions["B"].width = 24
    ws_staff.column_dimensions["C"].width = 22
    ws_staff.column_dimensions["I"].width = 30

    # ==========================================
    # SHEET 2: Daily Sign-in Sheet (Printable Register)
    # ==========================================
    ws_sign = wb.create_sheet(title="Printable_Sign_Sheet")
    ws_sign.views.sheetView[0].showGridLines = True

    ws_sign["A1"] = hostel_title
    ws_sign["A1"].font = font_title
    ws_sign["A2"] = f"DAILY HOUSEKEEPING ATTENDANCE & SIGNATURE REGISTER - {month_name.upper()} {year}"
    ws_sign["A2"].font = font_subtitle

    # Table Headers
    ws_sign.cell(row=4, column=1, value="Date").alignment = Alignment(horizontal="center", vertical="center")
    ws_sign.cell(row=4, column=2, value="Day").alignment = Alignment(horizontal="center", vertical="center")
    
    col_cur = 3
    for emp in staff_list:
        ws_sign.merge_cells(start_row=4, start_column=col_cur, end_row=4, end_column=col_cur+2)
        cell = ws_sign.cell(row=4, column=col_cur, value=f"{emp['name']} [{emp['location']}]")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.font = font_header
        cell.fill = fill_navy
        
        ws_sign.cell(row=5, column=col_cur, value="In Time / Sig").alignment = Alignment(horizontal="center")
        ws_sign.cell(row=5, column=col_cur+1, value="Out Time / Sig").alignment = Alignment(horizontal="center")
        ws_sign.cell(row=5, column=col_cur+2, value="Status (P/A/HD)").alignment = Alignment(horizontal="center")
        
        for c_offset in range(3):
            scell = ws_sign.cell(row=5, column=col_cur + c_offset)
            scell.font = font_bold
            scell.fill = fill_light_blue
            scell.border = dark_border
            
        col_cur += 3

    ws_sign.merge_cells(start_row=4, start_column=1, end_row=5, end_column=1)
    ws_sign.merge_cells(start_row=4, start_column=2, end_row=5, end_column=2)
    for r in (4, 5):
        for c in (1, 2):
            ws_sign.cell(row=r, column=c).fill = fill_navy
            ws_sign.cell(row=r, column=c).font = font_header
            ws_sign.cell(row=r, column=c).border = dark_border

    # Days rows
    for d in range(1, num_days + 1):
        r = 5 + d
        day_date = datetime.date(year, month, d)
        day_name = day_date.strftime("%a")
        is_sunday = day_date.weekday() == 6

        c_date = ws_sign.cell(row=r, column=1, value=d)
        c_date.alignment = Alignment(horizontal="center", vertical="center")
        c_day = ws_sign.cell(row=r, column=2, value=day_name)
        c_day.alignment = Alignment(horizontal="center", vertical="center")

        fill_to_use = fill_weekend if is_sunday else PatternFill(fill_type=None)
        c_date.fill = fill_to_use
        c_day.fill = fill_to_use
        c_date.border = dark_border
        c_day.border = dark_border

        c_idx = 3
        for _ in staff_list:
            c1 = ws_sign.cell(row=r, column=c_idx)
            c2 = ws_sign.cell(row=r, column=c_idx+1)
            c3 = ws_sign.cell(row=r, column=c_idx+2)
            for c in (c1, c2, c3):
                c.border = dark_border
                c.fill = fill_to_use
                c.alignment = Alignment(horizontal="center", vertical="center")
            if is_sunday:
                c1.value = "-"
                c2.value = "-"
                c3.value = "WO"
            c_idx += 3

    # Sign summary footer
    r_summary = 6 + num_days
    ws_sign.merge_cells(start_row=r_summary, start_column=1, end_row=r_summary, end_column=2)
    c_tot = ws_sign.cell(row=r_summary, column=1, value="Total Present Days")
    c_tot.font = font_bold
    c_tot.fill = fill_summary
    c_tot.alignment = Alignment(horizontal="center", vertical="center")
    c_tot.border = dark_border
    ws_sign.cell(row=r_summary, column=2).border = dark_border

    c_idx = 3
    for _ in staff_list:
        ws_sign.merge_cells(start_row=r_summary, start_column=c_idx, end_row=r_summary, end_column=c_idx+2)
        c_val = ws_sign.cell(row=r_summary, column=c_idx, value="________________ Days")
        c_val.font = font_bold
        c_val.fill = fill_summary
        c_val.alignment = Alignment(horizontal="center", vertical="center")
        for col_off in range(3):
            ws_sign.cell(row=r_summary, column=c_idx + col_off).border = dark_border
        c_idx += 3

    # Column dimensions for printable sign sheet
    ws_sign.column_dimensions["A"].width = 7
    ws_sign.column_dimensions["B"].width = 8
    for col_i in range(3, col_cur):
        ws_sign.column_dimensions[get_column_letter(col_i)].width = 14

    # ==========================================
    # SHEET 3: Monthly Attendance Matrix (Digital View)
    # ==========================================
    ws_matrix = wb.create_sheet(title="Attendance_Log")
    ws_matrix.views.sheetView[0].showGridLines = True

    ws_matrix["A1"] = hostel_title
    ws_matrix["A1"].font = font_title
    ws_matrix["A2"] = f"Monthly Attendance Log & Calculation Matrix - {month_name} {year}"
    ws_matrix["A2"].font = font_subtitle

    # Legend / Key
    ws_matrix["A4"] = "Legend:"
    ws_matrix["A4"].font = font_bold
    ws_matrix["B4"] = "P: Present (1.0)"
    ws_matrix["C4"] = "HD: Half Day (0.5)"
    ws_matrix["D4"] = "WO: Weekly Off (1.0 Paid)"
    ws_matrix["E4"] = "PL: Paid Leave (1.0 Paid)"
    ws_matrix["F4"] = "A: Absent / LOP (0.0)"
    for c in ["B4", "C4", "D4", "E4", "F4"]:
        ws_matrix[c].font = font_small

    headers_matrix = ["ID", "Staff Name", "Location"] + [str(d) for d in range(1, num_days + 1)] + [
        "Present (P)", "Half Day (HD)", "Weekly Off (WO)", "Paid Leave (PL)", "Absent (A)", "Total Payable Days"
    ]

    for col_idx, h in enumerate(headers_matrix, 1):
        cell = ws_matrix.cell(row=6, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_navy
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = dark_border

    # Sub-header with day names (Mon, Tue...)
    ws_matrix.cell(row=7, column=1, value="")
    ws_matrix.cell(row=7, column=2, value="")
    ws_matrix.cell(row=7, column=3, value="Day:")
    for d in range(1, num_days + 1):
        day_date = datetime.date(year, month, d)
        day_name = day_date.strftime("%a")
        cell = ws_matrix.cell(row=7, column=3 + d, value=day_name[:2])
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.font = font_small
        if day_date.weekday() == 6:
            cell.fill = fill_weekend
        else:
            cell.fill = fill_light_blue
        cell.border = thin_border

    for c in range(1, 4):
        ws_matrix.cell(row=7, column=c).border = thin_border
        ws_matrix.cell(row=7, column=c).fill = fill_light_blue

    for c in range(num_days + 4, len(headers_matrix) + 1):
        cell = ws_matrix.cell(row=7, column=c, value="")
        cell.border = thin_border
        cell.fill = fill_light_blue

    # Employee Attendance Rows
    row_cur = 8
    for emp in staff_list:
        ws_matrix.cell(row=row_cur, column=1, value=emp["id"]).alignment = Alignment(horizontal="center")
        ws_matrix.cell(row=row_cur, column=2, value=emp["name"]).font = font_bold
        ws_matrix.cell(row=row_cur, column=3, value=emp["location"])

        for c in range(1, 4):
            ws_matrix.cell(row=row_cur, column=c).border = thin_border

        for d in range(1, num_days + 1):
            day_date = datetime.date(year, month, d)
            col_pos = 3 + d
            cell = ws_matrix.cell(row=row_cur, column=col_pos)
            if day_date.weekday() == 6:
                cell.value = "WO"
                cell.fill = fill_weekend
            else:
                cell.value = "P"
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border

        # Formula columns
        first_day_col = get_column_letter(4)
        last_day_col = get_column_letter(3 + num_days)
        
        # Present Count
        col_p = 4 + num_days
        ws_matrix.cell(row=row_cur, column=col_p, value=f'=COUNTIF({first_day_col}{row_cur}:{last_day_col}{row_cur}, "P")').alignment = Alignment(horizontal="center")
        
        # Half Day Count
        col_hd = col_p + 1
        ws_matrix.cell(row=row_cur, column=col_hd, value=f'=COUNTIF({first_day_col}{row_cur}:{last_day_col}{row_cur}, "HD")').alignment = Alignment(horizontal="center")

        # Weekly Off Count
        col_wo = col_hd + 1
        ws_matrix.cell(row=row_cur, column=col_wo, value=f'=COUNTIF({first_day_col}{row_cur}:{last_day_col}{row_cur}, "WO")').alignment = Alignment(horizontal="center")

        # Paid Leave Count
        col_pl = col_wo + 1
        ws_matrix.cell(row=row_cur, column=col_pl, value=f'=COUNTIF({first_day_col}{row_cur}:{last_day_col}{row_cur}, "PL")').alignment = Alignment(horizontal="center")

        # Absent Count
        col_a = col_pl + 1
        ws_matrix.cell(row=row_cur, column=col_a, value=f'=COUNTIF({first_day_col}{row_cur}:{last_day_col}{row_cur}, "A")').alignment = Alignment(horizontal="center")

        # Total Payable Days = Present + (0.5 * HD) + WO + PL
        col_tot = col_a + 1
        let_p = get_column_letter(col_p)
        let_hd = get_column_letter(col_hd)
        let_wo = get_column_letter(col_wo)
        let_pl = get_column_letter(col_pl)
        
        c_payable = ws_matrix.cell(row=row_cur, column=col_tot, value=f'={let_p}{row_cur} + ({let_hd}{row_cur}*0.5) + {let_wo}{row_cur} + {let_pl}{row_cur}')
        c_payable.alignment = Alignment(horizontal="center")
        c_payable.font = font_bold
        c_payable.fill = fill_summary

        for c_summary in range(col_p, col_tot + 1):
            ws_matrix.cell(row=row_cur, column=c_summary).border = dark_border

        row_cur += 1

    # Column widths for Attendance Matrix
    ws_matrix.column_dimensions["A"].width = 10
    ws_matrix.column_dimensions["B"].width = 22
    ws_matrix.column_dimensions["C"].width = 20
    for d in range(1, num_days + 1):
        ws_matrix.column_dimensions[get_column_letter(3 + d)].width = 4.5
    for c in range(4 + num_days, len(headers_matrix) + 1):
        ws_matrix.column_dimensions[get_column_letter(c)].width = 14

    # ==========================================
    # SHEET 4: Monthly Remuneration & Payroll Calculation
    # ==========================================
    ws_pay = wb.create_sheet(title="Monthly_Remuneration")
    ws_pay.views.sheetView[0].showGridLines = True

    ws_pay["A1"] = hostel_title
    ws_pay["A1"].font = font_title
    ws_pay["A2"] = f"HOUSEKEEPING REMUNERATION & WAGE DISBURSEMENT REGISTER - {month_name.upper()} {year}"
    ws_pay["A2"].font = font_subtitle

    headers_pay = [
        "Staff ID", "Staff Name", "Location", "Monthly Base (₹)", f"Days in Month ({num_days})",
        "Payable Days", "Earned Basic Wage (₹)", "Special / Overtime Allowance (₹)",
        "Gross Remuneration (₹)", "Advance / Deductions (₹)", "NET PAYABLE WAGE (₹)",
        "Payment Mode", "Staff Signature / Date"
    ]

    for col_idx, h in enumerate(headers_pay, 1):
        cell = ws_pay.cell(row=4, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_navy
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = dark_border

    row_pay = 5
    for i, emp in enumerate(staff_list, start=1):
        staff_row_in_master = 4 + i
        staff_row_in_log = 7 + i
        
        # ID, Name, Location
        ws_pay.cell(row=row_pay, column=1, value=f"=Staff_Master!A{staff_row_in_master}").alignment = Alignment(horizontal="center")
        ws_pay.cell(row=row_pay, column=2, value=f"=Staff_Master!B{staff_row_in_master}").font = font_bold
        ws_pay.cell(row=row_pay, column=3, value=f"=Staff_Master!C{staff_row_in_master}")
        
        # Monthly Base Wage
        ws_pay.cell(row=row_pay, column=4, value=f"=Staff_Master!E{staff_row_in_master}").number_format = "₹#,##0.00"
        
        # Days in Month
        ws_pay.cell(row=row_pay, column=5, value=num_days).alignment = Alignment(horizontal="center")
        
        # Payable Days from Attendance Log
        payable_days_col_letter = get_column_letter(9 + num_days)
        ws_pay.cell(row=row_pay, column=6, value=f"=Attendance_Log!{payable_days_col_letter}{staff_row_in_log}").alignment = Alignment(horizontal="center")
        ws_pay.cell(row=row_pay, column=6).font = font_bold

        # Earned Basic Wage = (Monthly Base / Days in Month) * Payable Days
        ws_pay.cell(row=row_pay, column=7, value=f"=(D{row_pay}/E{row_pay})*F{row_pay}").number_format = "₹#,##0.00"

        # Overtime / Allowance (Default 0)
        ws_pay.cell(row=row_pay, column=8, value=0).number_format = "₹#,##0.00"

        # Gross Remuneration = Earned Basic + Overtime/Allowance
        ws_pay.cell(row=row_pay, column=9, value=f"=G{row_pay}+H{row_pay}").number_format = "₹#,##0.00"
        ws_pay.cell(row=row_pay, column=9).font = font_bold

        # Advance / Deductions (Default 0)
        ws_pay.cell(row=row_pay, column=10, value=0).number_format = "₹#,##0.00"

        # NET PAYABLE WAGE = Gross - Deductions
        c_net = ws_pay.cell(row=row_pay, column=11, value=f"=I{row_pay}-J{row_pay}")
        c_net.number_format = "₹#,##0.00"
        c_net.font = Font(name="Calibri", size=11, bold=True, color="006100")
        c_net.fill = fill_summary

        # Payment Mode
        ws_pay.cell(row=row_pay, column=12, value=f"=Staff_Master!H{staff_row_in_master}").alignment = Alignment(horizontal="center")

        # Signature placeholder
        ws_pay.cell(row=row_pay, column=13, value="").border = thin_border

        for col_c in range(1, 14):
            ws_pay.cell(row=row_pay, column=col_c).border = dark_border

        row_pay += 1

    # Totals Row
    ws_pay.merge_cells(start_row=row_pay, start_column=1, end_row=row_pay, end_column=3)
    c_tot_label = ws_pay.cell(row=row_pay, column=1, value="TOTAL REMUNERATION")
    c_tot_label.font = font_bold
    c_tot_label.alignment = Alignment(horizontal="center", vertical="center")
    c_tot_label.fill = fill_light_blue

    for col_c in (4, 7, 8, 9, 10, 11):
        col_let = get_column_letter(col_c)
        c_sum = ws_pay.cell(row=row_pay, column=col_c, value=f"=SUM({col_let}5:{col_let}{row_pay-1})")
        c_sum.font = font_bold
        c_sum.number_format = "₹#,##0.00"
        c_sum.fill = fill_light_blue
        if col_c == 11:
            c_sum.font = font_kpi

    for col_c in range(1, 14):
        ws_pay.cell(row=row_pay, column=col_c).border = double_bottom_border

    # Column dimensions for Remuneration Sheet
    for col in ws_pay.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_pay.column_dimensions[col_letter].width = max(max_len + 4, 15)
    ws_pay.column_dimensions["A"].width = 10
    ws_pay.column_dimensions["B"].width = 22
    ws_pay.column_dimensions["C"].width = 20
    ws_pay.column_dimensions["K"].width = 20
    ws_pay.column_dimensions["M"].width = 24

    # Save workbook
    wb.save(output_path)
    print(f"Excel workbook successfully created at: {output_path}")


def generate_printable_html(month, year, hostel_title, staff_list, out_html_attendance, out_html_payslip):
    num_days = calendar.monthrange(year, month)[1]
    month_name = calendar.month_name[month]

    # 1. Printable Attendance Register HTML
    days_header_html = "".join([f"<th style='width: 28px; text-align: center;'>{d}</th>" for d in range(1, num_days + 1)])
    
    days_sub_html = ""
    for d in range(1, num_days + 1):
        day_date = datetime.date(year, month, d)
        day_name = day_date.strftime("%a")[:2]
        is_sun = day_date.weekday() == 6
        bg_style = "background-color: #ffebee; color: #c62828; font-weight: bold;" if is_sun else "background-color: #f5f5f5;"
        days_sub_html += f"<th style='text-align: center; font-size: 10px; {bg_style}'>{day_name}</th>"

    rows_html = ""
    for emp in staff_list:
        in_cells = ""
        out_cells = ""
        for d in range(1, num_days + 1):
            day_date = datetime.date(year, month, d)
            is_sun = day_date.weekday() == 6
            if is_sun:
                in_cells += "<td style='background-color: #ffebee; text-align: center; font-size: 10px; color: #c62828;'>OFF</td>"
                out_cells += "<td style='background-color: #ffebee; text-align: center; font-size: 10px; color: #c62828;'>OFF</td>"
            else:
                in_cells += "<td style='height: 28px;'></td>"
                out_cells += "<td style='height: 28px;'></td>"

        rows_html += f"""
        <tr>
            <td rowspan="2" style="font-weight: bold; text-align: center; vertical-align: middle; background-color: #fafafa;">{emp['id']}</td>
            <td rowspan="2" style="font-weight: bold; vertical-align: middle; background-color: #fafafa;">
                <div style="font-size: 13px; color: #0f172a;">{emp['name']}</div>
                <div style="font-size: 11px; color: #1e40af; font-weight: 600;">{emp['location']}</div>
            </td>
            <td style="font-size: 11px; text-align: center; font-weight: 600; background: #eef2f7;">Morning / In</td>
            {in_cells}
            <td rowspan="2" style="text-align: center; vertical-align: middle; font-weight: bold;"></td>
            <td rowspan="2" style="text-align: center; vertical-align: middle; font-weight: bold;"></td>
        </tr>
        <tr>
            <td style="font-size: 11px; text-align: center; font-weight: 600; background: #eef2f7;">Evening / Out</td>
            {out_cells}
        </tr>
        """

    html_attendance = f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<title>Housekeeping Attendance Register - {month_name} {year}</title>
<style>
    @page {{
        size: A4 landscape;
        margin: 10mm;
    }}
    body {{
        font-family: 'Segoe UI', Arial, sans-serif;
        margin: 0;
        padding: 10px;
        color: #333;
    }}
    .header {{
        text-align: center;
        margin-bottom: 15px;
        border-bottom: 2px solid #1f497d;
        padding-bottom: 8px;
    }}
    .header h1 {{
        margin: 0;
        font-size: 20px;
        color: #1f497d;
        text-transform: uppercase;
        letter-spacing: 1px;
    }}
    .header h2 {{
        margin: 4px 0 0 0;
        font-size: 14px;
        color: #555;
        font-weight: normal;
    }}
    .meta-bar {{
        display: flex;
        justify-content: space-between;
        margin-bottom: 10px;
        font-size: 12px;
        font-weight: 600;
    }}
    table {{
        width: 100%;
        border-collapse: collapse;
        font-size: 11px;
    }}
    th, td {{
        border: 1px solid #777;
        padding: 4px 3px;
    }}
    th {{
        background-color: #1f497d;
        color: white;
        font-size: 11px;
    }}
    .legend {{
        margin-top: 15px;
        font-size: 11px;
        display: flex;
        gap: 20px;
        background: #f8f9fa;
        padding: 6px 12px;
        border: 1px solid #ddd;
        border-radius: 4px;
    }}
    .signatures {{
        margin-top: 30px;
        display: flex;
        justify-content: space-between;
        font-size: 12px;
        font-weight: bold;
    }}
    .sig-box {{
        border-top: 1px dashed #333;
        width: 220px;
        text-align: center;
        padding-top: 6px;
    }}
    @media print {{
        .no-print {{ display: none; }}
    }}
</style>
</head>
<body>

<div class="no-print" style="margin-bottom: 15px; text-align: right;">
    <button onclick="window.print()" style="background: #1f497d; color: white; border: none; padding: 8px 16px; font-size: 14px; border-radius: 4px; cursor: pointer;">🖨️ Print Attendance Sheet (A4 Landscape)</button>
</div>

<div class="header">
    <h1>{hostel_title}</h1>
    <h2>Monthly Housekeeping Staff Attendance Register | <strong>{month_name.upper()} {year}</strong></h2>
</div>

<div class="meta-bar">
    <span>Locations Covered: MBA Hostel &bull; Sports Hostel &bull; Toulouse Arena</span>
    <span>Total Days in Month: {num_days} Days</span>
    <span>Duty Timing: General Shift (7:00 AM - 4:00 PM)</span>
</div>

<table>
    <thead>
        <tr>
            <th rowspan="2" style="width: 55px; text-align: center;">ID</th>
            <th rowspan="2" style="width: 170px; text-align: left;">Staff Name & Location</th>
            <th rowspan="2" style="width: 75px; text-align: center;">Shift</th>
            {days_header_html}
            <th rowspan="2" style="width: 50px; text-align: center;">Present Days</th>
            <th rowspan="2" style="width: 70px; text-align: center;">Supervisor Sign</th>
        </tr>
        <tr>
            {days_sub_html}
        </tr>
    </thead>
    <tbody>
        {rows_html}
    </tbody>
</table>

<div class="legend">
    <span><strong>Attendance Codes:</strong></span>
    <span><strong>P</strong> = Present (Full Day)</span>
    <span><strong>HD</strong> = Half Day</span>
    <span><strong>WO</strong> = Weekly Off</span>
    <span><strong>PL</strong> = Approved Paid Leave</span>
    <span><strong>A</strong> = Absent (Unpaid)</span>
</div>

</body>
</html>
"""
    with open(out_html_attendance, "w", encoding="utf-8") as f:
        f.write(html_attendance)

    # 2. Printable Remuneration Register & Individual Pay Slips HTML
    slips_html = ""
    summary_rows_html = ""
    total_gross = 0
    total_net = 0

    for emp in staff_list:
        monthly_wage = emp["monthly_wage"]
        payable_days = num_days
        per_day_rate = monthly_wage / num_days
        earned_wage = monthly_wage
        ot_allowance = 0
        deductions = 0
        net_wage = earned_wage + ot_allowance - deductions
        total_gross += earned_wage
        total_net += net_wage

        summary_rows_html += f"""
        <tr>
            <td style="text-align: center; font-weight: bold;">{emp['id']}</td>
            <td style="font-weight: bold;">{emp['name']}</td>
            <td style="color: #1e40af; font-weight: 600;">{emp['location']}</td>
            <td style="text-align: right;">₹{monthly_wage:,.2f}</td>
            <td style="text-align: center;">{num_days}</td>
            <td style="text-align: center; font-weight: bold; background: #f0f7ff;">{payable_days}</td>
            <td style="text-align: right; font-weight: bold;">₹{earned_wage:,.2f}</td>
            <td style="text-align: right;">₹{ot_allowance:,.2f}</td>
            <td style="text-align: right;">₹{deductions:,.2f}</td>
            <td style="text-align: right; font-weight: bold; color: #1b5e20; background: #e8f5e9;">₹{net_wage:,.2f}</td>
            <td style="text-align: center;">{emp['payment_mode']}</td>
            <td style="height: 35px;"></td>
        </tr>
        """

        slips_html += f"""
        <div class="payslip-card">
            <div class="slip-header">
                <h3>{hostel_title}</h3>
                <p>Housekeeping Remuneration Voucher | <strong>{month_name} {year}</strong></p>
            </div>
            
            <div class="slip-grid">
                <div><strong>Staff ID:</strong> {emp['id']}</div>
                <div><strong>Month / Year:</strong> {month_name} {year}</div>
                <div><strong>Staff Name:</strong> {emp['name']}</div>
                <div><strong>Assigned Location:</strong> {emp['location']}</div>
                <div><strong>Designation:</strong> {emp['role']}</div>
                <div><strong>Payable Days:</strong> {payable_days} / {num_days} Days</div>
                <div><strong>Payment Mode:</strong> {emp['payment_mode']}</div>
                <div><strong>Bank / UPI Info:</strong> {emp['account_details']}</div>
            </div>

            <table class="slip-calc-table">
                <thead>
                    <tr>
                        <th>Earnings Component</th>
                        <th style="text-align: right;">Amount (₹)</th>
                        <th>Deductions / Advance</th>
                        <th style="text-align: right;">Amount (₹)</th>
                    </tr>
                </thead>
                <tbody>
                    <tr>
                        <td>Monthly Base Wage</td>
                        <td style="text-align: right;">₹{monthly_wage:,.2f}</td>
                        <td>Salary Advance / Loan</td>
                        <td style="text-align: right;">₹0.00</td>
                    </tr>
                    <tr>
                        <td>Earned Basic Remuneration</td>
                        <td style="text-align: right; font-weight: bold;">₹{earned_wage:,.2f}</td>
                        <td>Loss of Pay / Absenteeism</td>
                        <td style="text-align: right;">₹0.00</td>
                    </tr>
                    <tr>
                        <td>Overtime / Special Allowance</td>
                        <td style="text-align: right;">₹{ot_allowance:,.2f}</td>
                        <td>Other Deductions</td>
                        <td style="text-align: right;">₹0.00</td>
                    </tr>
                    <tr class="slip-total-row">
                        <td><strong>Total Gross Earnings</strong></td>
                        <td style="text-align: right; font-weight: bold;">₹{earned_wage:,.2f}</td>
                        <td><strong>Total Deductions</strong></td>
                        <td style="text-align: right; font-weight: bold;">₹{deductions:,.2f}</td>
                    </tr>
                </tbody>
            </table>

            <div class="net-pay-box">
                <span>NET PAYABLE AMOUNT:</span>
                <span class="net-amount">₹{net_wage:,.2f}</span>
            </div>

            <div class="slip-signatures">
                <div>
                    <br><br>
                    <span>Staff Signature (Acknowledgement)</span>
                </div>
                <div>
                    <br><br>
                    <span>Hostel Warden / Authorised Signatory</span>
                </div>
            </div>
        </div>
        """

    html_remuneration = f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<title>Housekeeping Remuneration & Payslips - {month_name} {year}</title>
<style>
    @page {{
        size: A4 portrait;
        margin: 10mm;
    }}
    body {{
        font-family: 'Segoe UI', Arial, sans-serif;
        margin: 0;
        padding: 15px;
        color: #333;
        background-color: #fcfcfc;
    }}
    .header {{
        text-align: center;
        margin-bottom: 20px;
        border-bottom: 2px solid #1f497d;
        padding-bottom: 10px;
    }}
    .header h1 {{
        margin: 0;
        font-size: 22px;
        color: #1f497d;
        text-transform: uppercase;
    }}
    .header h2 {{
        margin: 5px 0 0 0;
        font-size: 15px;
        color: #555;
    }}
    table {{
        width: 100%;
        border-collapse: collapse;
        margin-bottom: 25px;
        font-size: 12px;
    }}
    th, td {{
        border: 1px solid #777;
        padding: 6px 8px;
    }}
    th {{
        background-color: #1f497d;
        color: white;
    }}
    .payslip-card {{
        border: 2px solid #1f497d;
        border-radius: 6px;
        padding: 15px;
        margin-bottom: 25px;
        background: white;
        page-break-inside: avoid;
    }}
    .slip-header {{
        text-align: center;
        border-bottom: 1px solid #ccc;
        padding-bottom: 6px;
        margin-bottom: 10px;
    }}
    .slip-header h3 {{
        margin: 0;
        color: #1f497d;
        font-size: 16px;
    }}
    .slip-header p {{
        margin: 3px 0 0 0;
        font-size: 12px;
        color: #555;
    }}
    .slip-grid {{
        display: grid;
        grid-template-columns: 1fr 1fr;
        gap: 6px 20px;
        font-size: 12px;
        margin-bottom: 12px;
        background: #f8f9fa;
        padding: 10px;
        border-radius: 4px;
    }}
    .slip-calc-table {{
        margin-bottom: 10px;
        font-size: 11px;
    }}
    .slip-calc-table th {{
        background-color: #4f81bd;
    }}
    .slip-total-row {{
        background: #f1f5f9;
    }}
    .net-pay-box {{
        display: flex;
        justify-content: space-between;
        align-items: center;
        background: #e8f5e9;
        border: 1px solid #81c784;
        padding: 10px 15px;
        border-radius: 4px;
        font-weight: bold;
        font-size: 14px;
        color: #1b5e20;
    }}
    .net-amount {{
        font-size: 18px;
    }}
    .slip-signatures {{
        display: flex;
        justify-content: space-between;
        margin-top: 25px;
        font-size: 11px;
        font-weight: 600;
        text-align: center;
    }}
    .slip-signatures div {{
        width: 220px;
        border-top: 1px dashed #444;
        padding-top: 4px;
    }}
    @media print {{
        .no-print {{ display: none; }}
        body {{ background: white; padding: 0; }}
    }}
</style>
</head>
<body>

<div class="no-print" style="margin-bottom: 15px; text-align: right;">
    <button onclick="window.print()" style="background: #1f497d; color: white; border: none; padding: 8px 16px; font-size: 14px; border-radius: 4px; cursor: pointer;">🖨️ Print Remuneration Sheet & Payslips</button>
</div>

<div class="header">
    <h1>{hostel_title}</h1>
    <h2>Monthly Housekeeping Wage Disbursement Register | <strong>{month_name.upper()} {year}</strong></h2>
</div>

<table>
    <thead>
        <tr>
            <th>ID</th>
            <th>Staff Name</th>
            <th>Location</th>
            <th>Monthly Base</th>
            <th>Days</th>
            <th>Payable</th>
            <th>Earned Wage</th>
            <th>OT / Allow.</th>
            <th>Deductions</th>
            <th>Net Payable</th>
            <th>Mode</th>
            <th>Recipient Signature</th>
        </tr>
    </thead>
    <tbody>
        {summary_rows_html}
        <tr style="font-weight: bold; background: #eef2f7;">
            <td colspan="3" style="text-align: center;">TOTAL DISBURSEMENT</td>
            <td style="text-align: right;">-</td>
            <td style="text-align: center;">-</td>
            <td style="text-align: center;">-</td>
            <td style="text-align: right;">₹{total_gross:,.2f}</td>
            <td style="text-align: right;">₹0.00</td>
            <td style="text-align: right;">₹0.00</td>
            <td style="text-align: right; color: #1b5e20; font-size: 13px; background: #d4edda;">₹{total_net:,.2f}</td>
            <td colspan="2"></td>
        </tr>
    </tbody>
</table>

<h2 style="color: #1f497d; font-size: 16px; margin: 30px 0 15px 0; border-bottom: 1px solid #1f497d; padding-bottom: 5px;">
    Individual Staff Remuneration Vouchers (Signed Acknowledgements)
</h2>

{slips_html}

</body>
</html>
"""
    with open(out_html_payslip, "w", encoding="utf-8") as f:
        f.write(html_remuneration)


if __name__ == "__main__":
    now = datetime.datetime.now()
    current_month = now.month
    current_year = now.year

    hostel_title = "HOSTEL & FACILITY HOUSEKEEPING MANAGEMENT"

    # Specific 3 housekeeping staff members
    staff_members = [
        {
            "id": "HK-01",
            "name": "Ms. Arputhamani",
            "location": "MBA Hostel",
            "role": "Housekeeper / MBA Hostel",
            "phone": "-",
            "monthly_wage": 11000,
            "weekly_off": "Sunday",
            "payment_mode": "Bank / Cash",
            "account_details": "MBA Hostel Desk"
        },
        {
            "id": "HK-02",
            "name": "Mr. R. Sakthivelan",
            "location": "Sports Hostel",
            "role": "Housekeeper / Sports Hostel",
            "phone": "-",
            "monthly_wage": 11000,
            "weekly_off": "Sunday",
            "payment_mode": "Bank / Cash",
            "account_details": "Sports Hostel Desk"
        },
        {
            "id": "HK-03",
            "name": "Ms. Revathi",
            "location": "Toulouse Arena",
            "role": "Housekeeper / Toulouse Arena",
            "phone": "-",
            "monthly_wage": 11000,
            "weekly_off": "Sunday",
            "payment_mode": "Bank / Cash",
            "account_details": "Toulouse Arena Desk"
        }
    ]

    base_dir = os.path.dirname(os.path.abspath(__file__))
    excel_path = os.path.join(base_dir, "Hostel_Housekeeping_Attendance_and_Payroll.xlsx")
    html_att_path = os.path.join(base_dir, "Printable_Attendance_Register.html")
    html_pay_path = os.path.join(base_dir, "Printable_Remuneration_and_Payslips.html")

    generate_excel_system(current_month, current_year, hostel_title, staff_members, excel_path)
    generate_printable_html(current_month, current_year, hostel_title, staff_members, html_att_path, html_pay_path)
    print("Files successfully generated for Ms. Arputhamani, Mr. R. Sakthivelan, and Ms. Revathi!")
